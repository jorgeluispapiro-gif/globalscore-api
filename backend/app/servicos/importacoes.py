"""Leitura, validação e confirmação transacional de CSV e XLSX."""

import csv
import hashlib
import io
import json
import re
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from flask import current_app
from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from app.extensoes import banco
from app.modelos import (
    Avaliacao,
    BaseReferencia,
    Entidade,
    GrupoComparavel,
    Importacao,
    Indicador,
    ItemAvaliacao,
    ItemPopulacaoReferencia,
    Observacao,
    PerfilImportacao,
    Projeto,
    ValorConsolidadoBase,
)
from app.modelos.entidades import agora_utc
from app.servicos.motor_percentil import calcular_percentil_inc
from app.servicos.perfis_importacao import (
    PerfilPrecisaRevisaoError,
    gerar_assinatura_estrutura,
    montar_fotografia_estrutura,
    validar_referencias_perfil,
)


FORMATOS_ACEITOS = {".csv": "CSV", ".xlsx": "XLSX"}
MESES = {
    "jan": 1, "janeiro": 1, "fev": 2, "fevereiro": 2, "mar": 3, "marco": 3,
    "março": 3, "abr": 4, "abril": 4, "mai": 5, "maio": 5, "jun": 6,
    "junho": 6, "jul": 7, "julho": 7, "ago": 8, "agosto": 8, "set": 9,
    "setembro": 9, "out": 10, "outubro": 10, "nov": 11, "novembro": 11,
    "dez": 12, "dezembro": 12,
}


class AlertasPendentesError(ValueError):
    """Indica que o gestor ainda não aceitou conscientemente os alertas."""


class DependenciasImportacaoError(ValueError):
    """Bloqueia a anulação quando resultados congelados usam o lote."""

    def __init__(self, bases, avaliacoes):
        self.bases = bases
        self.avaliacoes = avaliacoes
        super().__init__(
            "A importação possui observações usadas por bases ou avaliações materializadas."
        )


class PerfilIncompativelError(ValueError):
    """Impede aplicação quando a estrutura atual difere da estrutura do perfil."""


def _resumo_perfil_sugerido(perfil):
    return {"id": perfil.id, "nome": perfil.nome, "versao": perfil.versao}


def importacao_para_dict(importacao, incluir_resumo=True):
    """Serializa metadados sem revelar o caminho temporário do servidor."""

    resposta = {
        "id": importacao.id,
        "projeto_id": importacao.projeto_id,
        "nome_arquivo_original": importacao.nome_arquivo_original,
        "tipo_arquivo": importacao.tipo_arquivo,
        "hash_sha256": importacao.hash_sha256,
        "status": importacao.status,
        "aba_selecionada": importacao.aba_selecionada,
        "configuracao_leitura": _carregar_json(importacao.configuracao_leitura_json),
        "mapeamento": _carregar_json(importacao.mapeamento_json),
        "perfil_importacao_id": importacao.perfil_importacao_id,
        "quantidade_linhas_lidas": importacao.quantidade_linhas_lidas,
        "quantidade_linhas_validas": importacao.quantidade_linhas_validas,
        "quantidade_erros": importacao.quantidade_erros,
        "quantidade_alertas": importacao.quantidade_alertas,
        "criado_por": importacao.criado_por,
        "criado_em": importacao.criado_em.isoformat(),
        "validado_em": importacao.validado_em.isoformat() if importacao.validado_em else None,
        "concluido_em": importacao.concluido_em.isoformat() if importacao.concluido_em else None,
    }
    if incluir_resumo:
        resposta["resumo_erros"] = _carregar_json(importacao.resumo_erros_json) or []
        resposta["resumo_alertas"] = _carregar_json(importacao.resumo_alertas_json) or []
    return resposta


def _carregar_json(valor):
    return json.loads(valor) if valor else None


def _valor_preview(valor):
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    return valor


def receber_arquivo(projeto_id, arquivo, criado_por="sistema"):
    projeto = banco.session.get(Projeto, projeto_id)
    if projeto is None:
        raise ValueError("Projeto não encontrado.")
    if arquivo is None or not arquivo.filename:
        raise ValueError("O arquivo é obrigatório.")

    nome_seguro = secure_filename(arquivo.filename)
    extensao = Path(nome_seguro).suffix.lower()
    if extensao not in FORMATOS_ACEITOS:
        raise ValueError("O MVP aceita somente arquivos CSV e XLSX.")

    conteudo = arquivo.read()
    if not conteudo:
        raise ValueError("O arquivo enviado está vazio.")
    identificador = f"{uuid4().hex}{extensao}"
    caminho = current_app.config["DIRETORIO_IMPORTACOES_TEMPORARIAS"] / identificador
    caminho.write_bytes(conteudo)

    importacao = Importacao(
        projeto_id=projeto.id,
        nome_arquivo_original=nome_seguro,
        tipo_arquivo=FORMATOS_ACEITOS[extensao],
        hash_sha256=hashlib.sha256(conteudo).hexdigest(),
        status="ENVIADA",
        caminho_arquivo_temporario=str(caminho),
        criado_por=criado_por,
    )
    banco.session.add(importacao)
    try:
        preview = inspecionar_arquivo(importacao)
        importacao.status = "ANALISADA"
        banco.session.commit()
    except (UnicodeDecodeError, OSError, ValueError, csv.Error) as erro:
        banco.session.rollback()
        caminho.unlink(missing_ok=True)
        raise ValueError("Não foi possível ler o arquivo como CSV ou XLSX válido.") from erro
    except Exception as erro:
        banco.session.rollback()
        caminho.unlink(missing_ok=True)
        # openpyxl pode lançar exceções específicas para pacotes XLSX corrompidos.
        raise ValueError("Não foi possível ler o arquivo como CSV ou XLSX válido.") from erro
    return importacao, preview


def inspecionar_arquivo(importacao):
    caminho = Path(importacao.caminho_arquivo_temporario)
    if importacao.tipo_arquivo == "CSV":
        texto = caminho.read_text(encoding="utf-8-sig")
        amostra = texto[:8192]
        try:
            delimitador = csv.Sniffer().sniff(amostra, delimiters=",;\t|").delimiter
        except csv.Error:
            delimitador = None
        leitor = csv.reader(io.StringIO(texto), delimiter=delimitador or ",")
        linhas = [[_valor_preview(valor) for valor in linha] for _, linha in zip(range(10), leitor)]
        return {"tipo_arquivo": "CSV", "sugestao_delimitador": delimitador, "preview": linhas}

    bruto = load_workbook(caminho, read_only=True, data_only=False)
    try:
        abas = []
        for nome in bruto.sheetnames:
            planilha = bruto[nome]
            linhas = [
                [_valor_preview(celula.value) for celula in linha]
                for _, linha in zip(range(10), planilha.iter_rows())
            ]
            abas.append({"nome": nome, "preview": linhas})
        return {"tipo_arquivo": "XLSX", "abas": abas}
    finally:
        bruto.close()


def _combinar_cabecalhos(linhas, indices, linhas_cabecalho):
    nomes = []
    for indice in indices:
        partes = []
        for numero_linha in linhas_cabecalho:
            linha = linhas[numero_linha - 1] if numero_linha <= len(linhas) else []
            valor = linha[indice - 1] if indice <= len(linha) else None
            if valor not in (None, ""):
                partes.append(str(valor).strip())
        nomes.append(" / ".join(partes) or f"Coluna {indice}")
    return nomes


def _ler_linhas(importacao, configuracao, incluir_todas_colunas=False):
    """Lê a região confirmada e preserva números de linha e fórmulas sem cache."""

    inicio = int(configuracao.get("linha_inicial", 0))
    fim = configuracao.get("linha_final")
    fim = int(fim) if fim is not None else None
    cabecalhos = configuracao.get("linhas_cabecalho", [])
    indices = configuracao.get("colunas_utilizadas", [])
    if inicio < 1 or len(cabecalhos) not in {1, 2} or (
        not indices and not incluir_todas_colunas
    ):
        raise ValueError(
            "Informe linha_inicial, uma ou duas linhas_cabecalho e colunas_utilizadas."
        )
    indices = [int(indice) for indice in indices]
    caminho = Path(importacao.caminho_arquivo_temporario or "")
    if not caminho.exists():
        raise ValueError("O arquivo temporário desta importação não está mais disponível.")

    formulas_sem_valor = set()
    formatos_celulas = {}
    if importacao.tipo_arquivo == "CSV":
        delimitador = configuracao.get("delimitador")
        if not delimitador:
            raise ValueError("Confirme o delimitador do CSV na configuração de leitura.")
        texto = caminho.read_text(encoding=configuracao.get("codificacao", "utf-8-sig"))
        linhas = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
    else:
        aba = configuracao.get("aba")
        if not aba:
            raise ValueError("Confirme a aba do XLSX na configuração de leitura.")
        livro_bruto = load_workbook(caminho, read_only=True, data_only=False)
        livro_valores = load_workbook(caminho, read_only=True, data_only=True)
        try:
            if aba not in livro_bruto.sheetnames:
                raise ValueError("A aba selecionada não existe no arquivo XLSX.")
            bruto = list(livro_bruto[aba].iter_rows())
            valores = list(livro_valores[aba].iter_rows())
            linhas = []
            for numero, linha in enumerate(valores, start=1):
                valores_linha = [celula.value for celula in linha]
                linhas.append(valores_linha)
                for coluna, celula in enumerate(livro_bruto[aba][numero], start=1):
                    formatos_celulas[(numero, coluna)] = celula.number_format
                    if celula.data_type == "f" and valores_linha[coluna - 1] is None:
                        formulas_sem_valor.add((numero, coluna))
        finally:
            livro_bruto.close()
            livro_valores.close()

    if incluir_todas_colunas:
        numeros_cabecalho = [int(item) for item in cabecalhos]
        quantidade_colunas = max(
            (
                len(linhas[numero - 1])
                for numero in numeros_cabecalho
                if numero <= len(linhas)
            ),
            default=0,
        )
        if quantidade_colunas == 0:
            raise ValueError("Não foi possível identificar os cabeçalhos do arquivo.")
        indices = list(range(1, quantidade_colunas + 1))

    nomes = _combinar_cabecalhos(linhas, indices, [int(item) for item in cabecalhos])
    registros = []
    limite = min(fim or len(linhas), len(linhas))
    for numero in range(inicio, limite + 1):
        linha = linhas[numero - 1]
        valores = {
            indice: (linha[indice - 1] if indice <= len(linha) else None)
            for indice in indices
        }
        if all(valor in (None, "") for valor in valores.values()):
            continue
        registros.append({"linha": numero, "valores": valores})
    return nomes, indices, registros, formulas_sem_valor, formatos_celulas


def ler_fotografia_estrutura_completa(importacao, configuracao_leitura):
    """Lê todos os cabeçalhos físicos para que nenhuma coluna passe despercebida."""

    configuracao = dict(configuracao_leitura or {})
    nomes, indices, *_ = _ler_linhas(
        importacao, configuracao, incluir_todas_colunas=True
    )
    configuracao["colunas_utilizadas"] = indices
    cabecalhos = [
        {"indice_coluna": indice, "nome": nome}
        for indice, nome in zip(indices, nomes)
    ]
    return montar_fotografia_estrutura(importacao, cabecalhos, configuracao)


def _estrutura_atual_para_perfil(importacao, perfil):
    """Lê a estrutura atual com a mesma configuração estrutural do perfil."""

    configuracao = json.loads(perfil.configuracao_leitura_json or "{}")
    return ler_fotografia_estrutura_completa(importacao, configuracao)


def _assinatura_atual_para_perfil(importacao, perfil):
    """Recalcula a assinatura usando a leitura estrutural configurada no perfil."""

    return gerar_assinatura_estrutura(_estrutura_atual_para_perfil(importacao, perfil))


def _agrupar_cabecalhos_por_nome(estrutura):
    """Agrupa cabeçalhos normalizados para impedir associações ambíguas."""

    agrupados = {}
    for cabecalho in estrutura.get("cabecalhos", []):
        nome = cabecalho.get("nome_normalizado", "")
        agrupados.setdefault(nome, []).append(cabecalho)
    return agrupados


def _comparar_estruturas(estrutura_perfil, estrutura_atual):
    """Compara nomes exatos e únicos e descreve somente diferenças observáveis."""

    cabecalhos_perfil = estrutura_perfil.get("cabecalhos", [])
    cabecalhos_atuais = estrutura_atual.get("cabecalhos", [])
    por_nome_perfil = _agrupar_cabecalhos_por_nome(estrutura_perfil)
    por_nome_atual = _agrupar_cabecalhos_por_nome(estrutura_atual)
    nomes_seguros = {
        nome
        for nome, itens_perfil in por_nome_perfil.items()
        if len(itens_perfil) == 1 and len(por_nome_atual.get(nome, [])) == 1
    }
    correspondencias = {
        por_nome_perfil[nome][0]["indice_coluna"]: por_nome_atual[nome][0]
        for nome in nomes_seguros
    }

    diferencas = []
    for cabecalho in cabecalhos_atuais:
        if cabecalho.get("nome_normalizado") not in nomes_seguros:
            diferencas.append(
                {
                    "tipo": "COLUNA_NOVA",
                    "indice_coluna": cabecalho["indice_coluna"],
                    "nome": cabecalho["nome_original"],
                }
            )
    for cabecalho in cabecalhos_perfil:
        if cabecalho.get("nome_normalizado") not in nomes_seguros:
            diferencas.append(
                {
                    "tipo": "COLUNA_REMOVIDA",
                    "indice_coluna": cabecalho["indice_coluna"],
                    "nome": cabecalho["nome_original"],
                }
            )

    ordem_perfil = [
        item["nome_normalizado"]
        for item in cabecalhos_perfil
        if item.get("nome_normalizado") in nomes_seguros
    ]
    ordem_atual = [
        item["nome_normalizado"]
        for item in cabecalhos_atuais
        if item.get("nome_normalizado") in nomes_seguros
    ]
    if ordem_perfil != ordem_atual:
        diferencas.append(
            {
                "tipo": "ORDEM_ALTERADA",
                "ordem_perfil": ordem_perfil,
                "ordem_atual": ordem_atual,
            }
        )
    return correspondencias, diferencas


def _montar_mapeamento_sugerido(perfil, correspondencias):
    """Reposiciona apenas colunas com nome normalizado exato e único."""

    mapeamento_perfil = json.loads(perfil.mapeamento_json or "{}")
    sugerido = {
        chave: deepcopy(valor)
        for chave, valor in mapeamento_perfil.items()
        if chave != "colunas"
    }
    colunas = []
    for coluna in mapeamento_perfil.get("colunas", []):
        cabecalho_atual = correspondencias.get(coluna.get("indice_coluna"))
        if cabecalho_atual is None:
            continue
        coluna_sugerida = deepcopy(coluna)
        coluna_sugerida["indice_coluna"] = cabecalho_atual["indice_coluna"]
        colunas.append(coluna_sugerida)
    sugerido["colunas"] = sorted(colunas, key=lambda item: item["indice_coluna"])
    return sugerido


def _contar_ancoras_perfil(perfil, correspondencias):
    """Conta somente papéis conhecidos que possuem correspondência estrutural segura."""

    mapeamento = json.loads(perfil.mapeamento_json or "{}")
    estruturais = 0
    indicadores = 0
    for coluna in mapeamento.get("colunas", []):
        if coluna.get("indice_coluna") not in correspondencias:
            continue
        if coluna.get("papel") in {"CODIGO_ENTIDADE", "PERIODO"}:
            estruturais += 1
        elif coluna.get("papel") == "VALOR_INDICADOR":
            indicadores += 1
    return estruturais, indicadores


def _diferencas_referencias(perfil):
    """Transforma uma referência inválida em pendência explícita de revisão."""

    try:
        validar_referencias_perfil(perfil)
    except PerfilPrecisaRevisaoError as erro:
        return [{"tipo": "REFERENCIA_INVALIDA", "mensagem": str(erro)}]
    return []


def _analisar_perfil(importacao, perfil):
    """Produz uma análise determinística, sem inferir nomes aproximados."""

    estrutura_perfil = json.loads(perfil.estrutura_json or "{}")
    estrutura_atual = _estrutura_atual_para_perfil(importacao, perfil)
    assinatura_exata = gerar_assinatura_estrutura(estrutura_atual) == perfil.assinatura_estrutura
    correspondencias, diferencas = _comparar_estruturas(
        estrutura_perfil, estrutura_atual
    )
    diferencas_referencias = _diferencas_referencias(perfil)
    estruturais, indicadores = _contar_ancoras_perfil(perfil, correspondencias)
    ancora_suficiente = (
        (estruturais >= 1 and indicadores >= 1) or indicadores >= 2
    )
    if not assinatura_exata and not ancora_suficiente:
        return None

    todas_diferencas = diferencas + diferencas_referencias
    return {
        "perfil": perfil,
        "assinatura_exata": assinatura_exata,
        "referencias_validas": not diferencas_referencias,
        "diferencas": todas_diferencas,
        "mapeamento_sugerido": _montar_mapeamento_sugerido(
            perfil, correspondencias
        ),
        # A tupla ordena candidatos por evidência objetiva. Empate permanece ambíguo.
        "pontuacao": (
            int(assinatura_exata),
            estruturais + indicadores,
            indicadores,
            estruturais,
            len(correspondencias),
        ),
    }


def reconhecer_perfil_importacao(importacao, criado_por):
    """Reconhece estruturas exatas ou parcialmente compatíveis sem fuzzy matching."""

    caminho = Path(importacao.caminho_arquivo_temporario or "")
    if not caminho.is_file():
        raise ValueError("O arquivo temporário desta importação não está mais disponível.")

    candidatos = PerfilImportacao.query.filter_by(
        projeto_id=importacao.projeto_id,
        criado_por=criado_por,
        ativo=True,
        tipo_arquivo=importacao.tipo_arquivo,
    ).order_by(PerfilImportacao.id).all()
    exatos = []
    parciais = []
    for perfil in candidatos:
        try:
            analise = _analisar_perfil(importacao, perfil)
        except (json.JSONDecodeError, OSError, UnicodeDecodeError, ValueError):
            continue
        if analise is None:
            continue
        if analise["assinatura_exata"] and analise["referencias_validas"]:
            exatos.append(analise)
        else:
            parciais.append(analise)

    if len(exatos) == 1:
        return {
            "resultado": "COMPATIVEL",
            "perfil_sugerido": _resumo_perfil_sugerido(exatos[0]["perfil"]),
        }
    if len(exatos) > 1:
        return {
            "resultado": "AMBIGUO",
            "perfil_sugerido": None,
            "candidatos": [
                _resumo_perfil_sugerido(item["perfil"]) for item in exatos
            ],
        }
    if not parciais:
        return {"resultado": "INCOMPATIVEL", "perfil_sugerido": None}

    melhor_pontuacao = max(item["pontuacao"] for item in parciais)
    melhores = [
        item for item in parciais if item["pontuacao"] == melhor_pontuacao
    ]
    if len(melhores) > 1:
        return {
            "resultado": "AMBIGUO",
            "perfil_sugerido": None,
            "candidatos": [
                _resumo_perfil_sugerido(item["perfil"]) for item in melhores
            ],
        }

    melhor = melhores[0]
    return {
        "resultado": "COMPATIVEL_COM_DIFERENCAS",
        "perfil_sugerido": _resumo_perfil_sugerido(melhor["perfil"]),
        "diferencas": melhor["diferencas"],
        "configuracao_leitura_sugerida": json.loads(
            melhor["perfil"].configuracao_leitura_json or "{}"
        ),
        "mapeamento_sugerido": melhor["mapeamento_sugerido"],
    }


def aplicar_perfil_importacao(importacao, perfil):
    """Aplica somente configuração e mapeamento após nova conferência exata."""

    if perfil.projeto_id != importacao.projeto_id:
        raise PerfilIncompativelError("O perfil não pertence ao projeto da importação.")
    if not perfil.ativo:
        raise PerfilPrecisaRevisaoError("O perfil está inativo e precisa ser revisado.")
    if perfil.tipo_arquivo != importacao.tipo_arquivo:
        raise PerfilIncompativelError("O tipo de arquivo não corresponde ao perfil.")

    assinatura_atual = _assinatura_atual_para_perfil(importacao, perfil)
    if assinatura_atual != perfil.assinatura_estrutura:
        raise PerfilIncompativelError(
            "A estrutura atual não é exatamente compatível com o perfil."
        )
    validar_referencias_perfil(perfil)

    importacao.configuracao_leitura_json = perfil.configuracao_leitura_json
    importacao.mapeamento_json = perfil.mapeamento_json
    importacao.perfil_importacao_id = perfil.id
    configuracao = json.loads(perfil.configuracao_leitura_json or "{}")
    importacao.aba_selecionada = configuracao.get("aba")
    banco.session.commit()
    return importacao


def _problema(tipo, mensagem, linha=None, coluna=None, valor=None):
    return {
        "linha": linha,
        "coluna": coluna,
        "valor_original": _valor_preview(valor),
        "tipo": tipo,
        "mensagem": mensagem,
    }


def _estatisticas_iqr(valores):
    """Calcula limites extremos de Tukey quando há amostra suficiente."""

    if len(valores) < 5:
        return None
    q1 = calcular_percentil_inc(valores, 25)
    q3 = calcular_percentil_inc(valores, 75)
    mediana = calcular_percentil_inc(valores, 50)
    iqr = q3 - q1
    if iqr == 0:
        return None
    return {
        "q1": q1,
        "q3": q3,
        "mediana": mediana,
        "iqr": iqr,
        "limite_inferior": q1 - Decimal("3") * iqr,
        "limite_superior": q3 + Decimal("3") * iqr,
        "tamanho_populacao": len(valores),
    }


def _identificacao(referencia):
    return referencia.get("id") or referencia.get("codigo")


def _montar_alerta(item, estatisticas, tipo):
    return {
        "tipo": tipo,
        "linha": item["linha"],
        "coluna": item["coluna"],
        "entidade": _identificacao(item["entidade"]),
        "indicador": _identificacao(item["indicador"]),
        "periodo": item["periodo"],
        "valor": str(item["valor"]),
        "q1": str(estatisticas["q1"]),
        "q3": str(estatisticas["q3"]),
        "mediana": str(estatisticas["mediana"]),
        "iqr": str(estatisticas["iqr"]),
        "limite_inferior": str(estatisticas["limite_inferior"]),
        "limite_superior": str(estatisticas["limite_superior"]),
        "tamanho_populacao": estatisticas["tamanho_populacao"],
    }


def _valor_atipico(valor, estatisticas):
    return valor < estatisticas["limite_inferior"] or valor > estatisticas["limite_superior"]


def _detectar_alertas_lote(plano):
    grupos = {}
    for item in plano:
        chave = (item["indicador"]["tipo"], _identificacao(item["indicador"]))
        grupos.setdefault(chave, []).append(item)
    alertas = []
    for itens in grupos.values():
        estatisticas = _estatisticas_iqr([item["valor"] for item in itens])
        if estatisticas:
            alertas.extend(
                _montar_alerta(item, estatisticas, "VALOR_ATIPICO_LOTE")
                for item in itens
                if _valor_atipico(item["valor"], estatisticas)
            )
    return alertas


def _detectar_alertas_historicos(importacao, plano):
    """Compara somente indicadores existentes com o histórico do mesmo grupo."""

    cache = {}
    alertas = []
    for item in plano:
        if item["indicador"]["tipo"] != "EXISTENTE":
            continue
        if item["entidade"]["tipo"] == "EXISTENTE":
            entidade = banco.session.get(Entidade, item["entidade"]["id"])
            grupo_id = entidade.grupo_id
        else:
            grupo_id = item["entidade"]["grupo_id"]
        chave = (item["indicador"]["id"], grupo_id)
        if chave not in cache:
            valores = [
                observacao.valor
                for observacao in Observacao.query.join(Entidade).filter(
                    Observacao.projeto_id == importacao.projeto_id,
                    Observacao.indicador_id == chave[0],
                    Entidade.grupo_id == grupo_id,
                )
            ]
            cache[chave] = _estatisticas_iqr(valores)
        estatisticas = cache[chave]
        if estatisticas and _valor_atipico(item["valor"], estatisticas):
            alertas.append(_montar_alerta(item, estatisticas, "VALOR_ATIPICO_HISTORICO"))
    return alertas


def normalizar_periodo(valor, formato=None):
    if isinstance(valor, (datetime, date)):
        return f"{valor.year:04d}-{valor.month:02d}"
    texto = str(valor or "").strip().lower()
    if not texto:
        raise ValueError("Período ausente.")
    if formato == "AAAA-MM" and re.fullmatch(r"\d{4}-\d{2}", texto):
        ano, mes = map(int, texto.split("-"))
    elif formato == "MM/AAAA" and re.fullmatch(r"\d{2}/\d{4}", texto):
        mes, ano = map(int, texto.split("/"))
    elif formato == "MES/AAAA" and re.fullmatch(r"[^/]+/\d{4}", texto):
        mes_texto, ano_texto = texto.split("/")
        mes, ano = MESES.get(mes_texto), int(ano_texto)
        if mes is None:
            raise ValueError("Nome de mês inválido.")
    else:
        raise ValueError("O período não corresponde ao formato confirmado.")
    if mes < 1 or mes > 12:
        raise ValueError("Mês inválido.")
    return f"{ano:04d}-{mes:02d}"


def normalizar_numero(valor, configuracao, percentual_excel=False):
    if valor in (None, ""):
        return None
    if isinstance(valor, (int, float, Decimal)):
        numero = Decimal(str(valor))
        percentual_textual = False
    else:
        texto = str(valor).strip()
        if not texto:
            return None
        percentual_textual = "%" in texto
        texto = texto.replace("%", "")
        for simbolo in configuracao.get("simbolos_monetarios", ["R$"]):
            texto = texto.replace(simbolo, "")
        texto = texto.replace(" ", "")
        milhar = configuracao.get("separador_milhar")
        decimal = configuracao.get("separador_decimal", ".")
        if milhar:
            texto = texto.replace(milhar, "")
        if decimal != ".":
            texto = texto.replace(decimal, ".")
        try:
            numero = Decimal(texto)
        except InvalidOperation as erro:
            raise ValueError("Valor numérico inválido.") from erro
    if percentual_textual or percentual_excel:
        modo = configuracao.get("percentual_como")
        if modo not in {"NUMERO", "FRACAO"}:
            raise ValueError("Confirme se percentuais devem ser NUMERO ou FRACAO.")
        if percentual_textual and modo == "FRACAO":
            numero /= Decimal("100")
        elif percentual_excel and modo == "NUMERO":
            numero *= Decimal("100")
    return numero


def normalizar_codigo_entidade(valor, formato_excel=None, origem_xlsx=False):
    """Preserva códigos numéricos somente para formatos XLSX simples como 000."""

    if valor in (None, ""):
        return ""
    if not origem_xlsx or not isinstance(valor, (int, float, Decimal)):
        return str(valor).strip()
    numero = Decimal(str(valor))
    if numero != numero.to_integral_value():
        raise ValueError("Código numérico de entidade não pode possuir casas decimais.")
    inteiro = int(numero)
    if formato_excel and re.fullmatch(r"0+", formato_excel):
        return str(inteiro).zfill(len(formato_excel))
    if formato_excel in {None, "", "General"}:
        return str(inteiro)
    raise ValueError(
        "Formato complexo no código da entidade; converta a coluna para texto ou use apenas zeros."
    )


def _validar_dados_novo_indicador(dados, projeto_id, problemas, coluna):
    obrigatorios = {
        "codigo",
        "nome",
        "unidade_medida",
        "direcao",
        "peso_percentual",
        "obrigatorio",
        "participa_global_score",
    }
    ausentes = sorted(campo for campo in obrigatorios if dados.get(campo) in (None, ""))
    if ausentes:
        problemas.append(_problema("CAMPO_OBRIGATORIO_AUSENTE", f"Dados do novo indicador ausentes: {', '.join(ausentes)}.", coluna=coluna))
        return False
    if dados["direcao"] not in {"MAIOR_MELHOR", "MENOR_MELHOR"}:
        problemas.append(_problema("VALOR_INVALIDO", "Direção do indicador inválida.", coluna=coluna, valor=dados["direcao"]))
        return False
    try:
        peso = Decimal(str(dados["peso_percentual"]))
    except InvalidOperation:
        peso = Decimal("-1")
    if peso < 0 or peso > 100:
        problemas.append(_problema("VALOR_INVALIDO", "O peso do indicador deve estar entre 0 e 100.", coluna=coluna, valor=dados["peso_percentual"]))
        return False
    existente = Indicador.query.filter_by(projeto_id=projeto_id, codigo=str(dados["codigo"])).first()
    if existente:
        problemas.append(_problema("DUPLICIDADE", "Já existe indicador com o código informado.", coluna=coluna, valor=dados["codigo"]))
        return False
    return True


def _resolver_entidade(importacao, codigo, nome, decisoes, problemas, linha):
    codigo = str(codigo or "").strip()
    if not codigo:
        problemas.append(_problema("CAMPO_OBRIGATORIO_AUSENTE", "Código da entidade ausente.", linha=linha))
        return None
    existente = Entidade.query.filter_by(projeto_id=importacao.projeto_id, codigo=codigo).first()
    if existente:
        return {"tipo": "EXISTENTE", "id": existente.id, "codigo": codigo}
    decisao = decisoes.get(codigo)
    if not decisao:
        problemas.append(_problema("ENTIDADE_DESCONHECIDA", "Defina se a nova entidade será criada, associada ou ignorada.", linha=linha, valor=codigo))
        return None
    if decisao.get("acao") == "IGNORAR":
        return {"tipo": "IGNORAR", "codigo": codigo}
    if decisao.get("acao") == "ASSOCIAR":
        entidade = banco.session.get(Entidade, decisao.get("entidade_id"))
        if entidade is None or entidade.projeto_id != importacao.projeto_id:
            problemas.append(_problema("ENTIDADE_DESCONHECIDA", "A entidade escolhida para associação não pertence ao projeto.", linha=linha, valor=codigo))
            return None
        return {"tipo": "EXISTENTE", "id": entidade.id, "codigo": codigo}
    if decisao.get("acao") == "CRIAR":
        grupo = banco.session.get(GrupoComparavel, decisao.get("grupo_id"))
        if grupo is None or grupo.projeto_id != importacao.projeto_id:
            problemas.append(_problema("CAMPO_OBRIGATORIO_AUSENTE", "A nova entidade exige grupo_id válido do projeto.", linha=linha, valor=codigo))
            return None
        return {"tipo": "NOVA", "codigo": codigo, "nome": decisao.get("nome") or str(nome or codigo), "grupo_id": grupo.id}
    problemas.append(_problema("ENTIDADE_DESCONHECIDA", "Ação de entidade inválida.", linha=linha, valor=codigo))
    return None


def _resolver_indicador(importacao, decisao, problemas, coluna):
    if not decisao or decisao.get("acao") == "IGNORAR":
        return {"tipo": "IGNORAR"} if decisao else None
    if decisao.get("acao") in {"ASSOCIAR", "EXISTENTE"}:
        indicador = banco.session.get(Indicador, decisao.get("indicador_id"))
        if indicador is None or indicador.projeto_id != importacao.projeto_id:
            problemas.append(_problema("INDICADOR_NAO_MAPEADO", "Indicador existente inválido para o projeto.", coluna=coluna))
            return None
        return {"tipo": "EXISTENTE", "id": indicador.id}
    if decisao.get("acao") == "CRIAR":
        dados = decisao.get("dados", {})
        if _validar_dados_novo_indicador(dados, importacao.projeto_id, problemas, coluna):
            return {"tipo": "NOVO", "codigo": str(dados["codigo"]), "dados": dados}
        return None
    problemas.append(_problema("INDICADOR_NAO_MAPEADO", "Ação de indicador inválida.", coluna=coluna))
    return None


def _executar_validacao(importacao, payload):
    configuracao = payload.get("configuracao_leitura", {})
    mapeamento = payload.get("mapeamento", {})
    nomes, indices, registros, formulas_sem_valor, formatos_celulas = _ler_linhas(
        importacao, configuracao
    )
    mapa_colunas = {int(item["indice_coluna"]): item for item in mapeamento.get("colunas", [])}
    formato = mapeamento.get("formato")
    problemas = []
    plano = []
    entidades_novas = {}
    indicadores_novos = {}
    chaves = set()

    if formato not in {"LARGO", "LONGO"}:
        raise ValueError("Confirme o formato LARGO ou LONGO no mapeamento.")
    for indice in indices:
        if indice not in mapa_colunas:
            problemas.append(_problema("INDICADOR_NAO_MAPEADO", "Toda coluna utilizada deve possuir papel explícito.", coluna=indice))

    papeis = {item.get("papel"): indice for indice, item in mapa_colunas.items()}
    if "CODIGO_ENTIDADE" not in papeis:
        raise ValueError("Mapeie uma coluna como CODIGO_ENTIDADE.")
    if "PERIODO" not in papeis and not configuracao.get("periodo_padrao"):
        raise ValueError("Mapeie PERIODO ou informe periodo_padrao.")

    for registro in registros:
        linha, valores = registro["linha"], registro["valores"]
        coluna_codigo = papeis["CODIGO_ENTIDADE"]
        codigo_original = valores.get(coluna_codigo)
        try:
            codigo = normalizar_codigo_entidade(
                codigo_original,
                formatos_celulas.get((linha, coluna_codigo)),
                importacao.tipo_arquivo == "XLSX",
            )
        except ValueError as erro:
            problemas.append(
                _problema("VALOR_INVALIDO", str(erro), linha, coluna_codigo, codigo_original)
            )
            continue
        nome = valores.get(papeis.get("NOME_ENTIDADE")) if papeis.get("NOME_ENTIDADE") else None
        entidade_ref = _resolver_entidade(importacao, codigo, nome, mapeamento.get("decisoes_entidades", {}), problemas, linha)
        if not entidade_ref or entidade_ref["tipo"] == "IGNORAR":
            continue
        if entidade_ref["tipo"] == "NOVA":
            entidades_novas[entidade_ref["codigo"]] = entidade_ref
        valor_periodo = valores.get(papeis.get("PERIODO")) if papeis.get("PERIODO") else configuracao.get("periodo_padrao")
        try:
            periodo = normalizar_periodo(valor_periodo, configuracao.get("formato_periodo"))
        except ValueError as erro:
            problemas.append(_problema("PERIODO_INVALIDO", str(erro), linha, papeis.get("PERIODO"), valor_periodo))
            continue

        pares = []
        if formato == "LARGO":
            for indice, item in mapa_colunas.items():
                if item.get("papel") == "VALOR_INDICADOR":
                    pares.append((indice, valores.get(indice), item.get("indicador")))
        else:
            coluna_indicador, coluna_valor = papeis.get("INDICADOR"), papeis.get("VALOR")
            if not coluna_indicador or not coluna_valor:
                raise ValueError("O formato LONGO exige colunas INDICADOR e VALOR.")
            rotulo = str(valores.get(coluna_indicador) or "").strip()
            pares.append((coluna_valor, valores.get(coluna_valor), mapeamento.get("decisoes_indicadores", {}).get(rotulo)))

        for coluna, valor_original, decisao_indicador in pares:
            if (linha, coluna) in formulas_sem_valor:
                problemas.append(_problema("VALOR_INVALIDO", "A fórmula não possui valor calculado armazenado.", linha, coluna, valor_original))
                continue
            if valor_original in (None, ""):
                continue
            indicador_ref = _resolver_indicador(importacao, decisao_indicador, problemas, coluna)
            if not indicador_ref or indicador_ref["tipo"] == "IGNORAR":
                if not decisao_indicador:
                    problemas.append(_problema("INDICADOR_NAO_MAPEADO", "O indicador precisa ser associado, criado ou ignorado.", linha, coluna))
                continue
            if indicador_ref["tipo"] == "NOVO":
                indicadores_novos[indicador_ref["codigo"]] = indicador_ref
            try:
                formato_excel = formatos_celulas.get((linha, coluna), "")
                valor = normalizar_numero(
                    valor_original,
                    configuracao.get("formato_numerico", {}),
                    importacao.tipo_arquivo == "XLSX" and "%" in formato_excel,
                )
            except ValueError as erro:
                problemas.append(_problema("VALOR_INVALIDO", str(erro), linha, coluna, valor_original))
                continue
            chave_entidade = (entidade_ref["tipo"], entidade_ref.get("id") or entidade_ref["codigo"])
            chave_indicador = (indicador_ref["tipo"], indicador_ref.get("id") or indicador_ref["codigo"])
            chave = (chave_entidade, chave_indicador, periodo)
            if chave in chaves:
                problemas.append(_problema("DUPLICIDADE", "A mesma combinação aparece mais de uma vez no arquivo.", linha, coluna, valor_original))
                continue
            chaves.add(chave)
            if entidade_ref["tipo"] == "EXISTENTE" and indicador_ref["tipo"] == "EXISTENTE":
                if Observacao.query.filter_by(entidade_id=entidade_ref["id"], indicador_id=indicador_ref["id"], periodo=periodo).first():
                    problemas.append(_problema("DUPLICIDADE", "Já existe observação para entidade, indicador e período.", linha, coluna, valor_original))
                    continue
            plano.append({"linha": linha, "coluna": coluna, "entidade": entidade_ref, "indicador": indicador_ref, "periodo": periodo, "valor": valor})

    alertas = _detectar_alertas_lote(plano) + _detectar_alertas_historicos(importacao, plano)

    resumo = {
        "linhas_lidas": len(registros),
        "linhas_validas": len({item["linha"] for item in plano}),
        "linhas_com_problema": len({item["linha"] for item in problemas if item["linha"]}),
        "entidades_reconhecidas": len({item["entidade"].get("id") for item in plano if item["entidade"]["tipo"] == "EXISTENTE"}),
        "novas_entidades": list(entidades_novas.values()),
        "indicadores_associados": len({item["indicador"].get("id") for item in plano if item["indicador"]["tipo"] == "EXISTENTE"}),
        "novos_indicadores": [item["dados"] for item in indicadores_novos.values()],
        "observacoes_a_criar": len(plano),
        "quantidade_erros": len(problemas),
        "quantidade_alertas": len(alertas),
        "duplicidades": sum(item["tipo"] == "DUPLICIDADE" for item in problemas),
        "valores_invalidos": sum(item["tipo"] == "VALOR_INVALIDO" for item in problemas),
        "periodos_invalidos": sum(item["tipo"] == "PERIODO_INVALIDO" for item in problemas),
        "cabecalhos": [{"indice_coluna": indice, "nome": nome} for indice, nome in zip(indices, nomes)],
        "problemas": problemas,
        "erros": problemas,
        "alertas": alertas,
        "_plano": plano,
    }
    return resumo


def validar_importacao(importacao_id, payload):
    importacao = banco.session.get(Importacao, importacao_id)
    if importacao is None:
        raise ValueError("Importação não encontrada.")
    if importacao.status in {"CONCLUIDA", "ANULADA", "CANCELADA"}:
        raise ValueError("Esta importação não pode ser revalidada em seu estado atual.")
    resumo = _executar_validacao(importacao, payload)
    importacao.aba_selecionada = payload.get("configuracao_leitura", {}).get("aba")
    importacao.configuracao_leitura_json = json.dumps(payload.get("configuracao_leitura", {}), ensure_ascii=False)
    importacao.estrutura_json = json.dumps(
        montar_fotografia_estrutura(
            importacao,
            resumo["cabecalhos"],
            payload.get("configuracao_leitura", {}),
        ),
        ensure_ascii=False,
    )
    importacao.mapeamento_json = json.dumps(payload.get("mapeamento", {}), ensure_ascii=False)
    importacao.quantidade_linhas_lidas = resumo["linhas_lidas"]
    importacao.quantidade_linhas_validas = resumo["linhas_validas"]
    importacao.quantidade_erros = resumo["quantidade_erros"]
    importacao.resumo_erros_json = json.dumps(resumo["problemas"], ensure_ascii=False, default=str)
    importacao.quantidade_alertas = resumo["quantidade_alertas"]
    importacao.resumo_alertas_json = json.dumps(resumo["alertas"], ensure_ascii=False)
    importacao.validado_em = agora_utc()
    if resumo["problemas"]:
        importacao.status = "ANALISADA"
    elif resumo["alertas"]:
        importacao.status = "VALIDADA_COM_ALERTAS"
    else:
        importacao.status = "VALIDADA"
    banco.session.commit()
    resumo.pop("_plano")
    resumo["status"] = importacao.status
    return resumo


def persistir_observacao(**dados):
    """Ponto pequeno e testável da gravação atômica do lote."""

    banco.session.add(Observacao(**dados))


def confirmar_importacao(importacao_id, confirmar_alertas=False):
    importacao = banco.session.get(Importacao, importacao_id)
    if importacao is None:
        raise ValueError("Importação não encontrada.")
    if importacao.status not in {"VALIDADA", "VALIDADA_COM_ALERTAS"}:
        raise ValueError("A importação precisa estar validada e sem erros antes da confirmação.")
    if importacao.status == "VALIDADA_COM_ALERTAS" and confirmar_alertas is not True:
        raise AlertasPendentesError(
            "A importação possui alertas. Envie confirmar_alertas=true para aceitá-los."
        )
    payload = {
        "configuracao_leitura": _carregar_json(importacao.configuracao_leitura_json) or {},
        "mapeamento": _carregar_json(importacao.mapeamento_json) or {},
    }
    resumo = _executar_validacao(importacao, payload)
    if resumo["problemas"]:
        raise ValueError("A importação deixou de ser válida; execute a validação novamente.")
    if resumo["alertas"] and confirmar_alertas is not True:
        raise AlertasPendentesError(
            "A importação possui alertas. Envie confirmar_alertas=true para aceitá-los."
        )

    entidades_criadas = {}
    indicadores_criados = {}
    try:
        for item in resumo["_plano"]:
            referencia = item["entidade"]
            if referencia["tipo"] == "NOVA" and referencia["codigo"] not in entidades_criadas:
                entidade = Entidade(
                    projeto_id=importacao.projeto_id,
                    grupo_id=referencia["grupo_id"],
                    codigo=referencia["codigo"],
                    nome=referencia["nome"],
                )
                banco.session.add(entidade)
                banco.session.flush()
                entidades_criadas[referencia["codigo"]] = entidade.id
            referencia_indicador = item["indicador"]
            if referencia_indicador["tipo"] == "NOVO" and referencia_indicador["codigo"] not in indicadores_criados:
                dados = referencia_indicador["dados"]
                indicador = Indicador(
                    projeto_id=importacao.projeto_id,
                    codigo=dados["codigo"], nome=dados["nome"],
                    unidade_medida=dados["unidade_medida"], direcao=dados["direcao"],
                    peso_percentual=Decimal(str(dados["peso_percentual"])),
                    obrigatorio=dados.get("obrigatorio", True),
                    participa_global_score=dados.get("participa_global_score", True),
                )
                banco.session.add(indicador)
                banco.session.flush()
                indicadores_criados[referencia_indicador["codigo"]] = indicador.id

        for item in resumo["_plano"]:
            entidade_id = item["entidade"].get("id") or entidades_criadas[item["entidade"]["codigo"]]
            indicador_id = item["indicador"].get("id") or indicadores_criados[item["indicador"]["codigo"]]
            persistir_observacao(
                projeto_id=importacao.projeto_id,
                entidade_id=entidade_id,
                indicador_id=indicador_id,
                periodo=item["periodo"],
                valor=item["valor"],
                origem="IMPORTACAO",
                importacao_id=importacao.id,
                criado_por=importacao.criado_por,
            )
        importacao.status = "CONCLUIDA"
        importacao.concluido_em = agora_utc()
        banco.session.commit()
    except Exception:
        banco.session.rollback()
        importacao = banco.session.get(Importacao, importacao_id)
        importacao.status = "FALHA"
        banco.session.commit()
        raise

    _remover_arquivo_temporario(importacao)
    importacao.caminho_arquivo_temporario = None
    banco.session.commit()
    return importacao, len(resumo["_plano"])


def _remover_arquivo_temporario(importacao):
    if importacao.caminho_arquivo_temporario:
        Path(importacao.caminho_arquivo_temporario).unlink(missing_ok=True)


def _dependencias_materializadas(importacao):
    """Localiza bases e avaliações que realmente consumiram observações do lote."""

    bases = set()
    avaliacoes = set()
    observacoes = Observacao.query.filter_by(importacao_id=importacao.id).all()
    for observacao in observacoes:
        itens_diretos = ItemPopulacaoReferencia.query.join(
            BaseReferencia, BaseReferencia.id == ItemPopulacaoReferencia.base_referencia_id
        ).filter(
            ItemPopulacaoReferencia.entidade_id == observacao.entidade_id,
            ItemPopulacaoReferencia.indicador_id == observacao.indicador_id,
            ItemPopulacaoReferencia.periodo == observacao.periodo,
            BaseReferencia.status.in_(["PROCESSADA", "ATIVA", "SUBSTITUIDA"]),
        )
        bases.update(item.base_referencia_id for item in itens_diretos)

        consolidados = ValorConsolidadoBase.query.join(
            BaseReferencia, BaseReferencia.id == ValorConsolidadoBase.base_referencia_id
        ).filter(
            ValorConsolidadoBase.entidade_id == observacao.entidade_id,
            ValorConsolidadoBase.indicador_id == observacao.indicador_id,
            ValorConsolidadoBase.status_cobertura == "ELEGIVEL",
            BaseReferencia.periodo_inicial <= observacao.periodo,
            BaseReferencia.periodo_final >= observacao.periodo,
            BaseReferencia.status.in_(["PROCESSADA", "ATIVA", "SUBSTITUIDA"]),
        )
        for consolidado in consolidados:
            materializado = ItemPopulacaoReferencia.query.filter_by(
                base_referencia_id=consolidado.base_referencia_id,
                entidade_id=observacao.entidade_id,
                indicador_id=observacao.indicador_id,
                origem="MEDIA_ENTIDADE",
            ).first()
            if materializado:
                bases.add(consolidado.base_referencia_id)

        itens_avaliacao = ItemAvaliacao.query.join(Avaliacao).filter(
            Avaliacao.entidade_id == observacao.entidade_id,
            Avaliacao.periodo == observacao.periodo,
            Avaliacao.status == "CALCULADA",
            ItemAvaliacao.indicador_id == observacao.indicador_id,
            ItemAvaliacao.valor_observado.is_not(None),
        )
        avaliacoes.update(item.avaliacao_id for item in itens_avaliacao)
    return sorted(bases), sorted(avaliacoes)


def anular_importacao(importacao_id):
    importacao = banco.session.get(Importacao, importacao_id)
    if importacao is None:
        raise ValueError("Importação não encontrada.")
    if importacao.status in {"ENVIADA", "ANALISADA", "VALIDADA", "VALIDADA_COM_ALERTAS"}:
        _remover_arquivo_temporario(importacao)
        importacao.caminho_arquivo_temporario = None
        importacao.status = "CANCELADA"
        banco.session.commit()
        return importacao, 0
    if importacao.status != "CONCLUIDA":
        raise ValueError("A importação já está anulada, cancelada ou falhou.")

    bases, avaliacoes = _dependencias_materializadas(importacao)
    if bases or avaliacoes:
        raise DependenciasImportacaoError(bases, avaliacoes)
    try:
        quantidade = Observacao.query.filter_by(importacao_id=importacao.id).delete()
        importacao.status = "ANULADA"
        banco.session.commit()
        return importacao, quantidade
    except Exception:
        banco.session.rollback()
        raise
